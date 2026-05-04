"""Strategy skeleton generator.

Bootstraps the per-strategy workspace ``strategies/<name>/`` from a research
paper plus a one-line logic description. The output is the *starting point*
for the research phase: a folder containing

* ``README.md``         — strategy summary (paper extract + parsed logic)
* ``hypothesis.yaml``   — valid spec (``automation.spec_validator`` clean)
* ``strategy.py``       — Python skeleton compatible with the optimizer
                           (``signals(df, **params) -> dict``)
* ``optimization_grid.yaml`` — parameter ranges for ``optuna_optimizer``
* ``data_cache/``       — empty (parquet drop zone for fetched OHLC)
* ``results/``          — empty (CSV / report drop zone)
* ``mql5_translation/`` — empty (translated EA drop zone)

The hypothesis is intentionally *plausible-but-placeholder*: every field
validates against the StrategySpec schema today, but the trader is expected
to refine the spec before running the full pipeline.

Paper ingestion supports PDF (pdfplumber), DOCX (python-docx) and plain
text/markdown (no extra deps). All optional; missing extractor → README
falls back to a "(paper extract unavailable)" note instead of failing.

CLI::

    python automation/generate_strategy.py \\
        --paper research_papers/mean_reversion_1.pdf \\
        --name mean_rev_1 \\
        --logic "RSI<25 + BB lower entry, intraday Asian session 22-06 UTC, max 16 bars" \\
        --symbols EURUSD,USDJPY,EURCHF,GBPNZD,EURAUD \\
        --is 2018-01-01:2023-12-31 \\
        --oos 2024-01-01:2026-04-30
"""
from __future__ import annotations

import argparse
import logging
import re
import sys
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import yaml

logger = logging.getLogger(__name__)

ROOT = Path(__file__).resolve().parent.parent
STRATEGIES_DIR = ROOT / "strategies"
SPECS_DIR = ROOT / "strategy_specs"

PAPER_EXTRACT_CHARS = 2000


# ---------------------------------------------------------------------------
# Paper ingestion
# ---------------------------------------------------------------------------

def extract_paper_text(paper_path: Path, max_chars: int = PAPER_EXTRACT_CHARS) -> str:
    """Best-effort plain-text extraction. Returns at most ``max_chars`` chars.

    Supported: ``.pdf`` (pdfplumber), ``.docx`` (python-docx),
    ``.md`` / ``.txt`` / ``.rst`` (read directly). Missing extractor or read
    failure → empty string and a warning, never an exception.
    """
    if not paper_path.exists():
        logger.warning("Paper not found: %s", paper_path)
        return ""

    suffix = paper_path.suffix.lower()
    text = ""
    try:
        if suffix == ".pdf":
            text = _extract_pdf(paper_path)
        elif suffix == ".docx":
            text = _extract_docx(paper_path)
        elif suffix in (".md", ".txt", ".rst"):
            text = paper_path.read_text(encoding="utf-8", errors="replace")
        else:
            logger.warning(
                "Unsupported paper extension %r; reading as plain text", suffix
            )
            text = paper_path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:  # noqa: BLE001
        logger.warning("Failed to extract %s: %s", paper_path, exc)
        return ""

    text = re.sub(r"\s+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text[:max_chars]


def _extract_pdf(path: Path) -> str:
    try:
        import pdfplumber  # type: ignore
    except ImportError:
        logger.warning(
            "pdfplumber not installed — cannot extract %s. "
            "Install with: pip install pdfplumber", path,
        )
        return ""
    chunks: List[str] = []
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            try:
                chunk = page.extract_text() or ""
            except Exception:  # noqa: BLE001
                chunk = ""
            if chunk:
                chunks.append(chunk)
            if sum(len(c) for c in chunks) > PAPER_EXTRACT_CHARS * 2:
                break
    return "\n\n".join(chunks)


def _extract_docx(path: Path) -> str:
    try:
        import docx  # type: ignore  # python-docx
    except ImportError:
        logger.warning(
            "python-docx not installed — cannot extract %s. "
            "Install with: pip install python-docx", path,
        )
        return ""
    document = docx.Document(str(path))
    return "\n".join(p.text for p in document.paragraphs if p.text.strip())


# ---------------------------------------------------------------------------
# Logic-description heuristics
# ---------------------------------------------------------------------------

EDGE_KEYWORDS = [
    ("mean_reversion", ["mean rev", "mean-rev", " mr ", " mr,", "fade", "revert",
                         "oversold", "overbought", "bollinger", "bb lower",
                         "bb upper"]),
    ("breakout", ["breakout", "break out", "donchian", "channel break"]),
    ("trend_following", ["trend follow", "trend-follow", "follow trend",
                          "ma cross", "ema cross", "moving average cross"]),
    ("momentum", ["momentum", "roc ", "rate of change", "rsi above", "rsi > 50"]),
    ("volatility", ["volatility", "atr breakout", "vol regime", "vix"]),
    ("carry", ["carry", "swap", "interest rate diff"]),
    ("arbitrage", ["arbitrage", "stat arb", "pairs trade", "cointegration"]),
]

REGIME_KEYWORDS = [
    ("low_vol_sideways", ["asian", "low vol", "low-vol", "sideways", "range",
                           "ranging"]),
    ("trending_bull", ["bull", "uptrend"]),
    ("trending_bear", ["bear", "downtrend"]),
    ("high_vol_crisis", ["high vol", "high-vol", "crisis", "panic", "vix"]),
]

HOLDING_KEYWORDS = [
    ("intraday", ["intraday", "session", "close before", "close at end"]),
    ("position_5_30d", ["weekly", "monthly", "position", "swing trade long",
                         "30 day", "30-day"]),
    ("swing_1_5d", ["swing", "multi-day", "multi day", "overnight"]),
]


def _match_keyword(text: str, table: List[Tuple[str, List[str]]], default: str) -> str:
    low = " " + text.lower() + " "
    for label, needles in table:
        for needle in needles:
            if needle in low:
                return label
    return default


def infer_edge_source(logic: str) -> str:
    return _match_keyword(logic, EDGE_KEYWORDS, "mean_reversion")


def infer_regime(logic: str) -> str:
    return _match_keyword(logic, REGIME_KEYWORDS, "low_vol_sideways")


def infer_holding_period(logic: str) -> str:
    return _match_keyword(logic, HOLDING_KEYWORDS, "intraday")


# Match "22-06", "22:00-06:00", "22 to 06" UTC sessions.
_SESSION_RE = re.compile(
    r"(\d{1,2})(?::(\d{2}))?\s*(?:-|–|to|–|—|/)\s*(\d{1,2})(?::(\d{2}))?",
    re.IGNORECASE,
)


def parse_session(logic: str) -> Tuple[str, str, bool]:
    """Return ``(start_hhmm, end_hhmm, crosses_midnight)``.

    Falls back to the FX Asian-session default ``22:00 → 06:00`` (cross-midnight)
    if no time range is detected. Always 24h, zero-padded.
    """
    m = _SESSION_RE.search(logic)
    if not m:
        return ("22:00", "06:00", True)
    h1, m1, h2, m2 = m.groups()
    h1, h2 = int(h1), int(h2)
    m1 = int(m1) if m1 else 0
    m2 = int(m2) if m2 else 0
    if not (0 <= h1 <= 23 and 0 <= h2 <= 23):
        return ("22:00", "06:00", True)
    crosses = h1 > h2 or (h1 == h2 and m1 > m2)
    return (f"{h1:02d}:{m1:02d}", f"{h2:02d}:{m2:02d}", crosses)


_MAX_BARS_RE = re.compile(r"max\s+(\d+)\s+bars?", re.IGNORECASE)


def parse_max_holding_bars(logic: str, default: int = 16) -> int:
    m = _MAX_BARS_RE.search(logic)
    return int(m.group(1)) if m else default


# ---------------------------------------------------------------------------
# Strategy ID allocation
# ---------------------------------------------------------------------------

_STR_ID_RE = re.compile(r"STR_(\d{3,})_", re.IGNORECASE)


def next_strategy_id(name: str, search_dirs: Optional[List[Path]] = None) -> str:
    """Allocate ``STR_NNN_<sanitized_name>`` choosing NNN as max-existing+1.

    Scans existing per-strategy folders (``strategies/<id>/``) and any spec
    files in ``strategy_specs/``. The minimum number is 1; numbers are
    zero-padded to 3 digits to match the validator regex.
    """
    if search_dirs is None:
        search_dirs = [STRATEGIES_DIR, SPECS_DIR]
    found: List[int] = []
    for d in search_dirs:
        if not d.exists():
            continue
        for entry in d.rglob("*"):
            m = _STR_ID_RE.search(entry.name)
            if m:
                try:
                    found.append(int(m.group(1)))
                except ValueError:
                    pass
    next_n = (max(found) + 1) if found else 1
    sanitized = re.sub(r"[^a-z0-9_]+", "_", name.lower()).strip("_") or "strategy"
    return f"STR_{next_n:03d}_{sanitized}"


# ---------------------------------------------------------------------------
# Spec / grid / strategy.py builders
# ---------------------------------------------------------------------------

def build_hypothesis_spec(
    *,
    strategy_id: str,
    name: str,
    paper_path: Path,
    logic_description: str,
    symbols: List[str],
    is_period: Tuple[date, date],
    oos_period: Tuple[date, date],
    forward_period: Optional[Tuple[date, date]] = None,
    author: str = "scyf",
) -> Dict:
    """Build a dict that validates as ``StrategySpec``.

    Forward window defaults to one year starting the day after OOS ends so
    the validator's ``oos_end < forward_start`` ordering holds.
    """
    if forward_period is None:
        fwd_start = oos_period[1] + timedelta(days=1)
        forward_period = (fwd_start, fwd_start + timedelta(days=365))

    sess_start, sess_end, crosses_midnight = parse_session(logic_description)
    max_bars = parse_max_holding_bars(logic_description)
    holding = infer_holding_period(logic_description)
    edge = infer_edge_source(logic_description)
    regime = infer_regime(logic_description)

    rationale = (
        f"Auto-generated hypothesis from research paper '{paper_path.name}'.\n"
        f"Logic summary: {logic_description.strip()}\n"
        "Mechanism placeholder — refine before running the pipeline. The author "
        "should rewrite this paragraph with a 3+ line behavioral or "
        "microstructure justification of why the edge exists, why it is durable, "
        "and why the chosen universe is the right one to harvest it."
    )

    failure_modes = (
        "- Regime change away from the assumption above.\n"
        "- High-impact news during the trading session.\n"
        "- Liquidity holes (broker close, holidays).\n"
        "- Parameter drift over time (validate with walk-forward)."
    )

    spec: Dict = {
        "meta": {
            "strategy_id": strategy_id,
            "author": author,
            "created_date": date.today().isoformat(),
            "hypothesis_version": "1.0",
            "tags": [edge, "auto_generated", name],
        },
        "hypothesis": {
            "rationale": rationale,
            "expected_edge_source": edge,
            "expected_regime_performance": regime,
            "expected_failure_modes": failure_modes,
            "expected_holding_period": holding,
        },
        "universe": {
            "symbols": list(symbols),
            "timeframe": "M15",
            "trading_session_utc": {
                "start": sess_start,
                "end": sess_end,
                "crosses_midnight": crosses_midnight,
            },
            "trading_days": ["Mon", "Tue", "Wed", "Thu"],
            "news_blackout": {
                "enabled": True,
                "minutes_before": 15,
                "minutes_after": 30,
            },
        },
        "entry_rules": {
            "description": f"Auto-stub from logic: {logic_description.strip()}",
            "long": {
                "enabled": True,
                "conditions": [
                    {
                        "type": "indicator_threshold",
                        "indicator": "RSI",
                        "params": {"period": 14, "applied_price": "close"},
                        "operator": "<",
                        "value": 25,
                    },
                    {
                        "type": "indicator_threshold",
                        "indicator": "BollingerBands",
                        "params": {"period": 20, "deviations": 2.0},
                        "operator": "price_below_lower",
                    },
                ],
            },
            "short": {
                "enabled": True,
                "conditions": [
                    {
                        "type": "indicator_threshold",
                        "indicator": "RSI",
                        "params": {"period": 14, "applied_price": "close"},
                        "operator": ">",
                        "value": 75,
                    },
                    {
                        "type": "indicator_threshold",
                        "indicator": "BollingerBands",
                        "params": {"period": 20, "deviations": 2.0},
                        "operator": "price_above_upper",
                    },
                ],
            },
            "filter_rules": [
                {
                    "description": "ATR floor — skip dead-flat conditions",
                    "type": "indicator_threshold",
                    "indicator": "ATR",
                    "params": {"period": 14},
                    "operator": ">",
                    "reference": "ATR(14) average over last 100 bars * 0.5",
                },
            ],
        },
        "exit_rules": {
            "stop_loss": {"type": "atr_multiple", "value": 1.5, "atr_period": 14},
            "take_profit": {"type": "atr_multiple", "value": 2.0, "atr_period": 14},
            "time_exit": {"enabled": True, "max_holding_bars": max_bars},
            "trailing": {"enabled": False},
            "reverse_on_opposite_signal": False,
        },
        "risk": {
            "position_sizing": "risk_per_trade_pct",
            "risk_per_trade_pct": 0.5,
            "max_concurrent_positions": min(len(symbols), 5),
            "max_concurrent_per_symbol": 1,
            "max_daily_loss_pct": 2.0,
            "max_drawdown_pct_circuit_breaker": 10.0,
            "max_slippage_points": 20,
        },
        "backtest": {
            "is_period": {
                "start": is_period[0].isoformat(),
                "end": is_period[1].isoformat(),
            },
            "oos_period": {
                "start": oos_period[0].isoformat(),
                "end": oos_period[1].isoformat(),
            },
            "forward_period": {
                "start": forward_period[0].isoformat(),
                "end": forward_period[1].isoformat(),
            },
            "initial_balance": 10000,
            "currency": "USD",
            "leverage": 30,
            "modeling": "every_tick_real",
            "use_real_spread": True,
            "optimization": False,
        },
        "acceptance_criteria": {
            "is_min_sharpe": 0.8,
            "oos_min_sharpe": 0.6,
            "bootstrap_max_pvalue": 0.01,
            "max_drawdown_pct": 12.0,
            "min_trades": 100,
            "param_sensitivity_min_retained": 0.5,
            "pnl_decomposition": {"min_directional_pnl_pct": 60},
        },
        "orthogonality_target": {
            "target_correlation_with_existing_book": 0.2,
            "target_uncovered_dimension": f"{edge}_{regime}",
            "orthogonality_rationale": (
                f"Auto-generated stub. Author should expand into a full "
                f"argument explaining why this {edge} strategy on the chosen "
                f"symbols is independent of the existing book — different "
                f"regime, different style, different time-of-day, or different "
                f"asset class."
            ),
        },
        "notes": (
            f"Source paper: {paper_path}\n"
            f"Logic: {logic_description.strip()}"
        ),
    }
    return spec


def build_optimization_grid(logic_description: str) -> Dict:
    """Return a default Optuna grid keyed under ``parameters``.

    The defaults are deliberately wide to give the optimizer room. Edit
    ``optimization_grid.yaml`` after generation to tighten ranges once the
    strategy logic is concrete.
    """
    bars = parse_max_holding_bars(logic_description, default=16)
    grid: Dict = {
        "parameters": {
            "rsi_period": {"type": "int", "low": 7, "high": 21, "step": 1},
            "rsi_oversold": {"type": "int", "low": 15, "high": 35, "step": 1},
            "rsi_overbought": {"type": "int", "low": 65, "high": 85, "step": 1},
            "bb_period": {"type": "int", "low": 10, "high": 40, "step": 2},
            "bb_dev": {"type": "float", "low": 1.5, "high": 3.0, "step": 0.1},
            "sl_atr_mult": {"type": "float", "low": 1.0, "high": 3.0, "step": 0.1},
            "tp_atr_mult": {"type": "float", "low": 1.0, "high": 4.0, "step": 0.1},
            "time_exit_bars": {
                "type": "int",
                "low": max(4, bars // 2),
                "high": max(bars * 2, bars + 4),
                "step": 1,
            },
        },
    }
    return grid


STRATEGY_PY_TEMPLATE = '''"""Strategy: __NAME__
Generated from: __PAPER_PATH__
Logic: __LOGIC__

Single source of truth for the research phase. After optimization the top
performers are translated to MQL5 (see ``mql5_translation/``).

Contract
--------
``signals(df, **params)`` returns a dict with::

    signal_long       pd.Series[bool]     # True at bar close → entry at next bar OPEN
    signal_short      pd.Series[bool]
    sl_distance       pd.Series[float]    # SL distance in price units
    tp_distance       pd.Series[float]    # TP distance in price units
    session_mask      pd.Series[bool]     # optional, True inside trading window
    max_holding_bars  int                 # optional, time-stop in bars

This signature is what ``python_engine.optuna_optimizer`` and
``python_engine.vectorized_backtest`` call directly — do not rename.
"""
from __future__ import annotations

from typing import Dict

import pandas as pd

from python_engine.indicators_mql5 import (  # noqa: F401
    sma, ema, smma, rsi, atr, bollinger_bands, macd, stochastic,
)


# Default parameters — overridden by optimization_grid.yaml during search.
DEFAULT_PARAMS: Dict = __DEFAULT_PARAMS__


def signals(df: pd.DataFrame, **params) -> Dict[str, pd.Series]:
    """Generate entry / exit signals.

    Parameters
    ----------
    df
        OHLC + spread M1 (or higher) panel from ``MT5DataFetcher``. Index is a
        ``DatetimeIndex`` in broker time.
    **params
        Strategy parameters; defaults come from ``DEFAULT_PARAMS``.
    """
    p = {**DEFAULT_PARAMS, **params}

    # AI EDIT: implement strategy logic here.
    #
    # Example sketch (Asian-session BB + RSI mean reversion):
    #
    #   rsi_val = rsi(df["close"], period=p["rsi_period"])
    #   mid, upper, lower = bollinger_bands(
    #       df["close"], period=p["bb_period"], deviations=p["bb_dev"],
    #   )
    #   atr_val = atr(df, period=14)
    #
    #   signal_long  = (rsi_val < p["rsi_oversold"])  & (df["close"] < lower)
    #   signal_short = (rsi_val > p["rsi_overbought"]) & (df["close"] > upper)
    #
    #   sl_distance = atr_val * p["sl_atr_mult"]
    #   tp_distance = atr_val * p["tp_atr_mult"]
    #
    #   hours = df.index.hour
    #   session_mask = (hours >= 22) | (hours < 6)
    #
    #   return {
    #       "signal_long":      signal_long.fillna(False).astype(bool),
    #       "signal_short":     signal_short.fillna(False).astype(bool),
    #       "sl_distance":      sl_distance,
    #       "tp_distance":      tp_distance,
    #       "session_mask":     pd.Series(session_mask, index=df.index),
    #       "max_holding_bars": p["time_exit_bars"],
    #   }

    raise NotImplementedError(
        "Fill in signals() with the strategy logic before running the optimizer."
    )
'''


def render_strategy_py(name: str, paper_path: Path, logic_description: str,
                       default_params: Dict) -> str:
    """Substitute placeholders without using str.format (user content may
    contain ``{`` / ``}``).
    """
    out = STRATEGY_PY_TEMPLATE
    out = out.replace("__NAME__", name)
    out = out.replace("__PAPER_PATH__", str(paper_path))
    out = out.replace("__LOGIC__", logic_description.strip())
    out = out.replace("__DEFAULT_PARAMS__", repr(default_params))
    return out


def render_readme(name: str, strategy_id: str, paper_path: Path,
                  logic_description: str, paper_extract: str,
                  symbols: List[str],
                  is_period: Tuple[date, date],
                  oos_period: Tuple[date, date]) -> str:
    extract_block = (
        paper_extract if paper_extract else "_(paper extract unavailable)_"
    )
    return (
        f"# Strategy: {name}\n\n"
        f"**ID:** `{strategy_id}`\n"
        f"**Source paper:** `{paper_path}`\n"
        f"**Generated:** {datetime.utcnow().isoformat(timespec='seconds')}Z\n\n"
        "## Hypothesis\n"
        f"{logic_description.strip()}\n\n"
        f"**Symbols:** {', '.join(symbols)}\n"
        f"**IS:** {is_period[0]} → {is_period[1]}\n"
        f"**OOS:** {oos_period[0]} → {oos_period[1]}\n\n"
        "## Status\n"
        "- [ ] Strategy code written\n"
        "- [ ] Optimization completed\n"
        "- [ ] Top-10 validated\n"
        "- [ ] MQL5 translation\n"
        "- [ ] MT5 validation backtest\n"
        "- [ ] Forward demo 30d\n\n"
        "## Top Performer\n"
        "_Filled after optimization._\n\n"
        "## Paper extract\n"
        "_First ~2000 chars of the source paper, for in-folder reference._\n\n"
        "```\n"
        f"{extract_block}\n"
        "```\n"
    )


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def generate_strategy_skeleton(
    paper_path: Path,
    name: str,
    logic_description: str,
    symbols: List[str],
    is_period: Tuple[date, date],
    oos_period: Tuple[date, date],
    *,
    forward_period: Optional[Tuple[date, date]] = None,
    author: str = "scyf",
    base_dir: Path = STRATEGIES_DIR,
    overwrite: bool = False,
) -> Path:
    """Create ``strategies/<name>/`` with all skeleton files.

    Returns the path of the created strategy folder. Raises ``FileExistsError``
    if it already exists and ``overwrite=False``.
    """
    if not symbols:
        raise ValueError("symbols list is empty")
    if is_period[0] >= is_period[1]:
        raise ValueError(f"is_period start must precede end: {is_period}")
    if oos_period[0] >= oos_period[1]:
        raise ValueError(f"oos_period start must precede end: {oos_period}")
    if is_period[1] >= oos_period[0]:
        raise ValueError(
            f"is_period.end {is_period[1]} must be < oos_period.start {oos_period[0]}"
        )

    base_dir.mkdir(parents=True, exist_ok=True)
    strategy_dir = base_dir / name
    if strategy_dir.exists():
        if not overwrite:
            raise FileExistsError(
                f"Strategy folder already exists: {strategy_dir}. "
                "Pass overwrite=True (or --overwrite) to recreate."
            )
    else:
        strategy_dir.mkdir(parents=True)

    for sub in ("data_cache", "results", "mql5_translation"):
        sub_path = strategy_dir / sub
        sub_path.mkdir(exist_ok=True)
        (sub_path / ".gitkeep").touch()

    strategy_id = next_strategy_id(name)
    paper_extract = extract_paper_text(paper_path)

    spec_dict = build_hypothesis_spec(
        strategy_id=strategy_id,
        name=name,
        paper_path=paper_path,
        logic_description=logic_description,
        symbols=symbols,
        is_period=is_period,
        oos_period=oos_period,
        forward_period=forward_period,
        author=author,
    )
    _validate_spec_dict(spec_dict)

    grid = build_optimization_grid(logic_description)

    (strategy_dir / "hypothesis.yaml").write_text(
        yaml.safe_dump(spec_dict, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    (strategy_dir / "optimization_grid.yaml").write_text(
        yaml.safe_dump(grid, sort_keys=False), encoding="utf-8",
    )
    (strategy_dir / "strategy.py").write_text(
        render_strategy_py(
            name=name,
            paper_path=paper_path,
            logic_description=logic_description,
            default_params=_default_params_from_grid(grid),
        ),
        encoding="utf-8",
    )
    (strategy_dir / "README.md").write_text(
        render_readme(
            name=name,
            strategy_id=strategy_id,
            paper_path=paper_path,
            logic_description=logic_description,
            paper_extract=paper_extract,
            symbols=symbols,
            is_period=is_period,
            oos_period=oos_period,
        ),
        encoding="utf-8",
    )

    logger.info("Generated strategy skeleton: %s", strategy_dir)
    return strategy_dir


def _validate_spec_dict(spec_dict: Dict) -> None:
    """Run the StrategySpec validator on the generated dict.

    Imported lazily so this module is usable in environments where pydantic
    isn't installed (e.g. doc-only rendering).
    """
    try:
        from automation.spec_validator import StrategySpec  # noqa: WPS433
    except Exception as exc:  # noqa: BLE001
        logger.warning("Skipping spec validation (%s)", exc)
        return
    StrategySpec(**spec_dict)


def _default_params_from_grid(grid: Dict) -> Dict:
    """Pick the midpoint of every range as the default param value."""
    out: Dict = {}
    for name, spec in grid.get("parameters", {}).items():
        ptype = spec.get("type", "float")
        if ptype == "categorical":
            out[name] = spec["choices"][0]
        else:
            lo, hi = spec["low"], spec["high"]
            mid = (lo + hi) / 2.0
            out[name] = int(round(mid)) if ptype == "int" else round(float(mid), 4)
    return out


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def _parse_period(s: str) -> Tuple[date, date]:
    try:
        a, b = s.split(":")
    except ValueError as exc:
        raise argparse.ArgumentTypeError(
            f"Period must be 'YYYY-MM-DD:YYYY-MM-DD', got {s!r}"
        ) from exc
    return (
        datetime.strptime(a.strip(), "%Y-%m-%d").date(),
        datetime.strptime(b.strip(), "%Y-%m-%d").date(),
    )


def _build_arg_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(
        prog="python automation/generate_strategy.py",
        description=(
            "Generate a strategies/<name>/ skeleton (README, hypothesis.yaml, "
            "strategy.py, optimization_grid.yaml) from a research paper."
        ),
    )
    p.add_argument("--paper", required=True, type=Path,
                   help="Path to the source paper (.pdf, .docx, .md, .txt).")
    p.add_argument("--name", required=True,
                   help="Folder name (snake_case), e.g. 'mean_rev_1'.")
    p.add_argument("--logic", required=True,
                   help="One-line plain-English logic summary.")
    p.add_argument("--symbols", required=True,
                   help="Comma-separated symbols, e.g. EURUSD,USDJPY,EURCHF.")
    p.add_argument("--is", dest="is_period", required=True, type=_parse_period,
                   help="In-sample 'YYYY-MM-DD:YYYY-MM-DD'.")
    p.add_argument("--oos", dest="oos_period", required=True, type=_parse_period,
                   help="Out-of-sample 'YYYY-MM-DD:YYYY-MM-DD'.")
    p.add_argument("--forward", dest="forward_period", default=None,
                   type=_parse_period,
                   help="Forward 'YYYY-MM-DD:YYYY-MM-DD'. Defaults to OOS+1y.")
    p.add_argument("--author", default="scyf")
    p.add_argument("--base-dir", default=STRATEGIES_DIR, type=Path,
                   help="Where to create the strategies/ folder.")
    p.add_argument("--overwrite", action="store_true",
                   help="Replace files in an existing folder of this name.")
    p.add_argument("-v", "--verbose", action="store_true")
    return p


def main(argv: Optional[List[str]] = None) -> int:
    args = _build_arg_parser().parse_args(argv)
    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(asctime)s | %(levelname)s | %(message)s",
    )
    symbols = [s.strip() for s in args.symbols.split(",") if s.strip()]
    if not symbols:
        print("--symbols must contain at least one symbol", file=sys.stderr)
        return 2
    folder = generate_strategy_skeleton(
        paper_path=args.paper,
        name=args.name,
        logic_description=args.logic,
        symbols=symbols,
        is_period=args.is_period,
        oos_period=args.oos_period,
        forward_period=args.forward_period,
        author=args.author,
        base_dir=args.base_dir,
        overwrite=args.overwrite,
    )
    print(f"Created: {folder}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
